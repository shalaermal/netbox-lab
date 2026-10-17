"""Build docs/LAB20-netbox-import.xlsx: a small campus (2 core, 2 agg, 10 access)
as NetBox import sheets, one sheet per Import page, in import order."""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
SITE, N = "LAB20", 20
GROUP = "LAB20 VLANs"
CORE = [f"{SITE}-CORE-01", f"{SITE}-CORE-02"]
AGG = [f"{SITE}-AGG-01", f"{SITE}-AGG-02"]
ACC = [(f"{SITE}-F{f}-ACC-{k:02d}", f) for f in (1, 2) for k in range(1, 6)]
M8360 = "Aruba 8360-48Y6C Front-to-Back"
M6300 = "Aruba 6300M-48G-POE4-4SFP56"
VLANS = [(10, "Users"), (20, "Voice"), (30, "WiFi"), (99, "Mgmt"), (999, "Parking")]
TRUNK = "10,20,30,99"
MTU = 9198

F = "Arial"
HEAD = Font(name=F, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", start_color="1F4E78")
BODY = Font(name=F)
NOTE = Font(name=F, italic=True, color="595959")

wb = Workbook()


def sheet(title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h) or max(len(str(h)), *(len(str(r[i - 1])) for r in rows)) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(w, 45)
    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------- 01..05 site, VLANs, IPAM
sheet("01_site", ["name", "slug", "status", "description"],
      [(SITE, SITE.lower(), "active", "Test campus: 2 core, 2 agg, 10 access")])
sheet("02_locations", ["site", "name", "slug", "status"],
      [(SITE, "MDF", f"{SITE.lower()}-mdf", "active"),
       (SITE, "Floor 1", f"{SITE.lower()}-f1", "active"),
       (SITE, "Floor 2", f"{SITE.lower()}-f2", "active")])
sheet("03_vlan_group", ["name", "slug", "scope_type", "scope_name"],
      [(GROUP, f"{SITE.lower()}-vlans", "dcim.site", SITE)])
sheet("04_vlans", ["group", "vid", "name", "status", "description"],
      [(GROUP, v, n, "active", "Unused ports, no gateway" if v == 999 else "") for v, n in VLANS])
pfx = [(f"10.{N}.{v}.0/24", "", GROUP, v, "active", "dcim.site", SITE, n) for v, n in VLANS if v != 999]
pfx += [(f"10.{N}.251.0/24", "", "", "", "container", "dcim.site", SITE, "Core-agg links (/31)"),
        (f"10.{N}.252.0/24", "KA", "", "", "container", "dcim.site", SITE, "VSX keepalive (/31)"),
        (f"10.{N}.253.0/24", "mgmt", "", "", "active", "dcim.site", SITE, "Out-of-band management"),
        (f"10.{N}.254.0/24", "", "", "", "container", "dcim.site", SITE, "Loopbacks (/32)")]
sheet("05_prefixes", ["prefix", "vrf", "vlan_group", "vlan", "status", "scope_type", "scope_name",
                      "description"], pfx)

# ---------------------------------------------------------------- 06 devices
dev = [(n, "Core", "HPE", M8360, "Aruba AOS-CX", SITE, "MDF", "active") for n in CORE]
dev += [(n, "Agg", "HPE", M8360, "Aruba AOS-CX", SITE, "MDF", "active") for n in AGG]
dev += [(n, "Access", "HPE", M6300, "Aruba AOS-CX", SITE, f"Floor {f}", "active") for n, f in ACC]
sheet("06_devices", ["name", "role", "manufacturer", "device_type", "platform", "site", "location",
                     "status"], dev)

# ---------------------------------------------------------------- 07 new interfaces
new = []
for d in CORE + AGG:
    new.append((d, "lag256", "lag", "tagged-all", "", "", MTU, "", "VSX ISL"))
    new.append((d, "loopback0", "virtual", "", "", "", "", "", "Router ID"))
for i, (a, _) in enumerate(ACC, 1):
    for g in AGG:
        new.append((g, f"lag{i}", "lag", "tagged", GROUP, TRUNK, MTU, "", f"MC-LAG to {a}"))
for g in AGG:
    for v, n in VLANS:
        if v != 999:
            new.append((g, f"vlan{v}", "virtual", "", "", "", "", "", n))
for a, _ in ACC:
    new.append((a, "lag1", "lag", "tagged", GROUP, TRUNK, MTU, "", "Uplink to AGG VSX"))
    new.append((a, "vlan99", "virtual", "", "", "", "", "", "Management"))
sheet("07_interfaces_new", ["device", "name", "type", "mode", "vlan_group", "tagged_vlans", "mtu",
                            "vrf.name", "description"], new)

# ---------------------------------------------------------------- 08 existing ports (update by id)
upd = []
p2p = [(CORE[0], "1/1/49", AGG[0], "1/1/49"), (CORE[0], "1/1/50", AGG[1], "1/1/49"),
       (CORE[1], "1/1/49", AGG[0], "1/1/50"), (CORE[1], "1/1/50", AGG[1], "1/1/50")]
for d in CORE + AGG:
    upd.append((d, "mgmt", "", "", "", "", "", "", "mgmt", "true", "OOB management"))
    upd.append((d, "1/1/48", "", "", "", "", "", "", "KA", "true", "VSX keepalive"))
    for p in ("1/1/53", "1/1/54"):
        upd.append((d, p, "lag256", "", "", "", "", "", "", "true", "VSX ISL member"))
for c, cp, a, ap in p2p:
    upd.append((c, cp, "", "", "", "", "", MTU, "", "true", f"P2P to {a}"))
    upd.append((a, ap, "", "", "", "", "", MTU, "", "true", f"P2P to {c}"))
for i, (acc, _) in enumerate(ACC, 1):
    for k, g in enumerate(AGG):
        upd.append((g, f"1/1/{i}", f"lag{i}", "", "", "", "", "", "", "true", f"To {acc} 1/1/{49 + k}"))
        upd.append((acc, f"1/1/{49 + k}", "lag1", "", "", "", "", "", "", "true", f"To {g} 1/1/{i}"))
for acc, _ in ACC:
    for p in range(1, 49):
        if p <= 40:
            upd.append((acc, f"1/1/{p}", "", "tagged", GROUP, 10, 20, "", "", "true", "USER"))
        elif p <= 44:
            upd.append((acc, f"1/1/{p}", "", "access", GROUP, 30, "", "", "", "true", "AP"))
        else:
            upd.append((acc, f"1/1/{p}", "", "access", GROUP, 999, "", "", "", "false", "UNUSED"))
headers = ["id", "device", "name", "lag", "mode", "vlan_group", "untagged_vlan", "tagged_vlans",
           "mtu", "vrf.name", "enabled", "description"]
rows = [("",) + r[:7] + (r[7], r[8], r[9], r[10]) for r in upd]
ws = sheet("08_ports_update", headers, rows, widths={"id": 9})
for r in range(2, len(rows) + 2):
    ws.cell(r, 1).value = (f'=IFERROR(INDEX(_ids!$A:$A,MATCH(B{r}&"|"&C{r},_ids!$D:$D,0)),'
                           f'"MISSING")')
    ws.cell(r, 1).font = Font(name=F, color="000000")

# helper sheet: paste the interface export here
ids = wb.create_sheet("_ids")
ids.append(["ID", "Device", "Name", "key (formula, do not edit)"])
for c in ids[1]:
    c.font, c.fill = HEAD, HEAD_FILL
for r in range(2, 1502):
    ids.cell(r, 4).value = f'=IF(A{r}="","",B{r}&"|"&C{r})'
    ids.cell(r, 4).font = NOTE
for col, w in zip("ABCD", (10, 22, 14, 30)):
    ids.column_dimensions[col].width = w
ids.freeze_panes = "A2"
ids.cell(1, 6).value = ("Paste the interface export here (columns A-C: ID, Device, Name). "
                        "Column D fills itself; 08_ports_update reads the IDs from it.")
ids.cell(1, 6).font = Font(name=F, bold=True, color="C00000")

# ---------------------------------------------------------------- 09 cables
cab = []
for pair in (CORE, AGG):
    for p in ("1/1/53", "1/1/54", "1/1/48"):
        cab.append((pair[0], "dcim.interface", p, pair[1], "dcim.interface", p, "dac-passive",
                    "connected", "VSX keepalive" if p == "1/1/48" else "VSX ISL"))
for c, cp, a, ap in p2p:
    cab.append((c, "dcim.interface", cp, a, "dcim.interface", ap, "smf", "connected", "core-agg"))
for i, (acc, _) in enumerate(ACC, 1):
    for k, g in enumerate(AGG):
        cab.append((acc, "dcim.interface", f"1/1/{49 + k}", g, "dcim.interface", f"1/1/{i}",
                    "mmf-om4", "connected", "access uplink"))
sheet("09_cables", ["side_a_device", "side_a_type", "side_a_name", "side_b_device", "side_b_type",
                    "side_b_name", "type", "status", "label"], cab)

# ---------------------------------------------------------------- 10 IP addresses
ip = []
oob = {CORE[0]: 11, CORE[1]: 12, AGG[0]: 21, AGG[1]: 22}
lo = {CORE[0]: 1, CORE[1]: 2, AGG[0]: 11, AGG[1]: 12}
for d in CORE + AGG:
    ip.append((f"10.{N}.253.{oob[d]}/24", "mgmt", "active", "", d, "mgmt", "true", "OOB management"))
    ip.append((f"10.{N}.254.{lo[d]}/32", "", "active", "", d, "loopback0", "false", "Router ID"))
for k, d in enumerate(CORE + AGG):
    ip.append((f"10.{N}.252.{k}/31", "KA", "active", "", d, "1/1/48", "false", "VSX keepalive"))
for n, (c, cp, a, ap) in enumerate(p2p):
    ip.append((f"10.{N}.251.{2 * n}/31", "", "active", "", c, cp, "false", f"P2P to {a}"))
    ip.append((f"10.{N}.251.{2 * n + 1}/31", "", "active", "", a, ap, "false", f"P2P to {c}"))
for v, nm in VLANS:
    if v == 999:
        continue
    for k, g in enumerate(AGG):
        ip.append((f"10.{N}.{v}.{2 + k}/24", "", "active", "", g, f"vlan{v}", "false", nm))
        ip.append((f"10.{N}.{v}.1/24", "", "active", "vip", g, f"vlan{v}", "false",
                   f"{nm} gateway (VSX active-gateway)"))
for i, (acc, _) in enumerate(ACC, 1):
    ip.append((f"10.{N}.99.{100 + i}/24", "", "active", "", acc, "vlan99", "true", "Management"))
sheet("10_ip_addresses", ["address", "vrf", "status", "role", "device", "interface", "is_primary",
                          "description"], ip)

# ---------------------------------------------------------------- 11 config contexts (GUI)
std = {"dns_servers": [f"10.{N}.30.53"], "ntp_servers": [f"10.{N}.30.123"],
       "syslog_servers": [f"10.{N}.30.60"], "dhcp_servers": [f"10.{N}.30.67"],
       "tacacs_servers": [f"10.{N}.30.49"],
       "snmp": {"v3_user": "netmon", "location_prefix": SITE},
       "ospf": {"process": 1, "area": "0.0.0.0"},
       "spanning_tree": {"mode": "mstp", "config_name": SITE, "config_revision": 1},
       "active_gateway_mac": "12:00:00:00:01:00", "access_mgmt_vlan": 99,
       "timezone": "europe/berlin", "banner": "Authorized access only. All activity is logged.\n"}
ctx = wb.create_sheet("11_context_GUI")
ctx.column_dimensions["A"].width = 24
ctx.column_dimensions["B"].width = 34
ctx.column_dimensions["C"].width = 110
ctx.append(["Where", "Object", "JSON to paste"])
for c in ctx[1]:
    c.font, c.fill = HEAD, HEAD_FILL
ctx.append(["Customization > Config Contexts > Add (Name, Weight 1000, Sites = LAB20, Data)",
            f"standards-{SITE.lower()}", json.dumps(std)])
macs = {"core": "02:00:20:00:00:01", "agg": "02:00:20:00:00:02"}
for pair, kind in ((CORE, "core"), (AGG, "agg")):
    for k, d in enumerate(pair):
        ctx.append(["Device > Edit > Local config context data", d,
                    json.dumps({"vsx": {"role": "primary" if k == 0 else "secondary",
                                        "system_mac": macs[kind], "peer": pair[1 - k]}})])
for row in ctx.iter_rows(min_row=2):
    for c in row:
        c.font = BODY
        c.alignment = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------- README (first sheet)
rd = wb.active
rd.title = "README"
rd.column_dimensions["A"].width = 6
rd.column_dimensions["B"].width = 26
rd.column_dimensions["C"].width = 44
rd.column_dimensions["D"].width = 70
lines = [
    ("title", "LAB20: si plotesohet NetBox-i me import (2 core, 2 agg, 10 access)"),
    ("p", "Cdo flete eshte NJE faqe Import ne NetBox. Importohen me radhe, sepse cdo objekt i referohet "
          "objekteve te krijuara me pare (pajisja i duhet site-i, kablli i duhen portat...)."),
    ("p", "Si importohet nje flete: selekto te gjitha qelizat me te dhena (Ctrl+A ne flete), Ctrl+C, "
          "ne NetBox hap faqen e Import-it, ngjit te kutia 'Data', Submit. Rreshti 1 (emrat e kolonave) "
          "duhet te jete gjithmone brenda."),
    ("p", "Kerkohet qe ne NetBox te ekzistojne tashme: manufacturer HPE, device types Aruba 8360 dhe "
          "6300M, rolet Core/Agg/Access, platforma 'Aruba AOS-CX' me config template 'Aruba CX', "
          "VRF-te 'mgmt' dhe 'KA'. Te gjitha krijohen nga 'python -m nblab setup-gui' + Build site LAB01."),
    ("h", ""),
    ("row", "#", "Fleta", "Ku ne NetBox", "Cfare krijon / shenime"),
    ("row", 1, "01_site", "Organization > Sites > Import", "Site-i LAB20."),
    ("row", 2, "02_locations", "Organization > Locations > Import", "MDF, Floor 1, Floor 2 (dhomat/katet)."),
    ("row", 3, "03_vlan_group", "IPAM > VLAN Groups > Import",
     "Grupi i VLAN-eve, i lidhur me site-in (scope). Pa kete, VLAN-et nuk dalin ne config."),
    ("row", 4, "04_vlans", "IPAM > VLANs > Import", "5 VLAN: 10 Users, 20 Voice, 30 WiFi, 99 Mgmt, 999 Parking."),
    ("row", 5, "05_prefixes", "IPAM > Prefixes > Import",
     "Subnet-et 10.20.x.0/24, secili i lidhur me VLAN-in; rangjet infra (p2p, keepalive, OOB, loopback)."),
    ("row", 6, "06_devices", "Devices > Devices > Import",
     "14 switch-a. NetBox krijon VETE portat 1/1/1...1/1/54 + mgmt nga modeli (device type)."),
    ("row", 7, "07_interfaces_new", "Devices > Interfaces > Import",
     "Interface-at qe NUK vijne nga modeli: LAG-et (lag1, lag256), SVI-te (vlan10...), loopback0."),
    ("row", 8, "08_ports_update", "Devices > Interfaces > Import",
     "Portat ekzistuese (te krijuara ne hapin 6) perditesohen: LAG-u, VLAN-et, VRF, MTU, pershkrimi. "
     "Kerkon kolonen 'id': shiko hapin 8a."),
    ("row", "8a", "_ids", "Devices > Interfaces",
     "Filtro Site = LAB20 > Configure Table: vetem ID, Device, Name (me kete radhe) > Export > "
     "Current view. Ngjit kolonat ne _ids!A:C. Kolona 'id' te 08 mbushet vete (nese del MISSING, "
     "mungon ai port ne export)."),
    ("row", 9, "09_cables", "Devices > Cables > Import", "30 kabllo: ISL, keepalive, core-agg, access uplinks."),
    ("row", 10, "10_ip_addresses", "IPAM > IP Addresses > Import",
     "IP-te. role=vip eshte gateway-i (VSX active-gateway), i njejti ne te dy AGG. is_primary=true = IP e "
     "menaxhimit te pajisjes."),
    ("row", 11, "11_context_GUI", "Me dore ne GUI (pa import)",
     "Standardet (NTP, DNS, TACACS...) dhe VSX per 4 pajisjet: kopjo JSON-in nga kjo flete."),
    ("row", 12, "-", "Customization > Scripts > Validate site", "Zgjidh LAB20. Duhet: 0 errors."),
    ("row", 13, "-", "Device > Render Config", "Config-i Aruba CX i cdo switch-i."),
    ("h", ""),
    ("p", "Si ta perdoresh per nje site tjeter: zevendeso LAB20 / 10.20 me site-in e ri ne te gjitha "
          "fletet (Ctrl+H), ndrysho pajisjet ne 06 dhe lidhjet ne 08/09. Rregulli: cdo port qe permendet "
          "ne 08, 09 ose 10 duhet te ekzistoje (nga modeli ose nga 07)."),
    ("p", "Kujdes: te interface-at kolona quhet 'vrf.name', sepse NetBox-i e kerkon VRF-ne e interface-it "
          "me RD; me '.name' i thua ta kerkoje me emer. E njejta sintakse (fusha.atributi) vlen per cdo kolone."),
    ("p", "Legjenda: rreshti 1 i cdo flete = emrat e fushave te NetBox-it, mos i ndrysho. Kolonat bosh "
          "lejohen. Ne 08 kolona 'id' eshte formule: mos e shkruaj me dore."),
]
r = 1
for item in lines:
    kind = item[0]
    if kind == "title":
        rd.cell(r, 2, item[1]).font = Font(name=F, bold=True, size=14)
    elif kind == "p":
        rd.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c = rd.cell(r, 2, item[1])
        c.font = BODY
        c.alignment = Alignment(wrap_text=True, vertical="top")
        rd.row_dimensions[r].height = 32
    elif kind == "row":
        vals = item[1:]
        hdr = vals[0] == "#"
        for j, v in enumerate(vals):
            c = rd.cell(r, 1 + j, v)
            c.font = HEAD if hdr else BODY
            if hdr:
                c.fill = HEAD_FILL
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if not hdr and len(str(vals[3])) > 70:
            rd.row_dimensions[r].height = 15 * (len(str(vals[3])) // 65 + 1)
    r += 1

wb.save(OUT)
print("saved", OUT, {ws.title: ws.max_row - 1 for ws in wb.worksheets})
